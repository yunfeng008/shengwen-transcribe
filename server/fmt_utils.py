#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 result.json 重新生成 txt / srt / docx（说话人合并等编辑操作后调用）。"""

import json
import subprocess
from collections import Counter
from pathlib import Path


def fmt_mmss(sec):
    sec = int(sec)
    h, sec = divmod(sec, 3600)
    m, s = divmod(sec, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def fmt_srt_time(sec):
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3600 * 1000)
    m, ms = divmod(ms, 60 * 1000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def regenerate_outputs(job_dir, data):
    """job_dir: 任务目录；data: result.json 的 dict（已改过 sentences）。"""
    job_dir = Path(job_dir)
    sentences = data["sentences"]
    names = data.get("speaker_names", {})

    def display(spk_display):
        return names.get(str(spk_display)) or f"讲话人{spk_display}"

    # txt
    with open(job_dir / "result.txt", "w", encoding="utf-8") as f:
        for s in sentences:
            f.write(f"{display(s['spk_display'])}  {fmt_mmss(s['start'])}  {s['text']}\n")

    # srt
    with open(job_dir / "result.srt", "w", encoding="utf-8") as f:
        for i, s in enumerate(sentences, 1):
            f.write(f"{i}\n{fmt_srt_time(s['start'])} --> {fmt_srt_time(s['end'])}\n"
                    f"[{display(s['spk_display'])}] {s['text']}\n\n")

    # docx
    _write_docx(job_dir / "result.docx", data)


def _write_docx(docx_path, data):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = data.get("title", "转写结果")
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    st = data.get("stats", {})
    names = data.get("speaker_names", {})
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = []
    if st.get("audio_minutes"):
        parts.append(f"时长 {fmt_mmss(st['audio_minutes'] * 60)}")
    parts.append(f"讲话人 {st.get('n_speakers', '?')} 位")
    parts.append(f"共 {st.get('n_segments', len(data['sentences']))} 句 / "
                 f"{st.get('n_chars', sum(len(s['text']) for s in data['sentences']))} 字")
    parts.append("已人工合并说话人" if st.get("merged") else "")
    r = meta.add_run("　|　".join(p for p in parts if p))
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # 讲话人名单（有真名时列出）
    if names:
        roster = doc.add_paragraph()
        roster.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = roster.add_run("　".join(f"{v}（讲话人{k}）" for k, v in sorted(names.items(), key=lambda x: int(x[0]))))
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(0x1a, 0x5c, 0xb0)

    def display(spk_display):
        return names.get(str(spk_display)) or f"讲话人{spk_display}"

    cur_spk, buf, t_start = None, [], None
    def flush():
        nonlocal cur_spk, buf, t_start
        if cur_spk is not None and buf:
            p = doc.add_paragraph()
            lab = p.add_run(f"{display(cur_spk)}  [{fmt_mmss(t_start)}]  ")
            lab.bold = True
            lab.font.color.rgb = RGBColor(0x1a, 0x5c, 0xb0)
            p.add_run("".join(buf))
        cur_spk, buf, t_start = None, [], None

    for s in data["sentences"]:
        if s["spk_display"] != cur_spk:
            flush()
            cur_spk, t_start = s["spk_display"], s["start"]
        buf.append(s["text"])
    flush()
    doc.save(str(docx_path))

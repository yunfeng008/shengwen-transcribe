#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""可选的 LLM 深度纠错：读取 result.json，用大模型修正同音错字和术语。

配置从 runs/config.json 读取（用户在 Web 设置页填）：
  {"llm_api_key": "...", "llm_api_base": "https://api.deepseek.com",
   "llm_model": "deepseek-chat", "domain": "religion"}
"""

import json
import os
import time
import urllib.request
import urllib.error
from pathlib import Path

APP_ROOT = Path(__file__).parent.parent.resolve()
CONFIG_FILE = APP_ROOT / "runs" / "config.json"

WINDOW_SEC = 180  # 每次给 LLM 的上下文窗口

DOMAIN_PROMPTS = {
    "religion": "这是一段信仰团体（教会）录音的转写。内容涉及圣经书卷、神学概念、人名。请基于这个背景修正识别错误。",
    "meeting": "这是一段中文会议录音的转写。",
    "lecture": "这是一段授课/讲座录音的转写。",
    "general": "这是一段中文语音转写。",
}


def load_config():
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text())
        except json.JSONDecodeError:
            pass
    return {}


def save_config(cfg):
    CONFIG_FILE.parent.mkdir(exist_ok=True)
    CONFIG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2))


def call_llm(api_base, api_key, model, messages, timeout=180, max_retries=3):
    url = f"{api_base.rstrip('/')}/chat/completions"
    payload = {"model": model, "messages": messages,
               "temperature": 0.2, "response_format": {"type": "json_object"}}
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(url, data=data, headers={
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}"}, method="POST")
    last_err = None
    for attempt in range(max_retries):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                result = json.loads(resp.read().decode("utf-8"))
            return result["choices"][0]["message"]["content"]
        except urllib.error.HTTPError as e:
            last_err = f"HTTP {e.code}: {e.read().decode('utf-8', 'ignore')[:150]}"
            if e.code in (429, 500, 502, 503):
                time.sleep(2 ** attempt * 3)
                continue
            raise
        except Exception as e:
            last_err = str(e)
            time.sleep(2 ** attempt * 3)
    raise RuntimeError(f"LLM 调用失败：{last_err}")


def build_messages(seg_text, domain, hotwords):
    system = (
        "你是中文语音转写（ASR）的后处理校正专家。\n\n"
        f"## 录音背景\n{DOMAIN_PROMPTS.get(domain, DOMAIN_PROMPTS['general'])}\n\n"
        "## 修正原则\n"
        "- 修正同音错字和术语（如『题摩态』→『提摩太』、『我阿门』→『我们』）\n"
        "- 保持口语风格和语气词，不润色、不总结、不删减原意\n"
        "- 保持每行『讲话人N  时间  内容』格式，行数不变\n"
        "- 整行外语噪声（韩文/日文碎片）内容清空\n\n"
        '## 输出严格 JSON：{"corrected_text": "..."}\n'
    )
    user = ""
    if hotwords:
        user += f"## 参考热词\n{', '.join(hotwords)}\n\n"
    user += f"## 请修正以下转写文本\n\n{seg_text}"
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def fmt_mmss(sec):
    sec = int(sec)
    h, sec = divmod(sec, 3600)
    m, s = divmod(sec, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def polish_result(job_dir, result, cb=None, hotwords=None, domain=None):
    """对 pipeline 结果做 LLM 纠错，生成 result_polished.* 文件。"""
    cfg = load_config()
    api_key = cfg.get("llm_api_key") or os.environ.get("LLM_API_KEY", "")
    api_base = cfg.get("llm_api_base") or os.environ.get("LLM_API_BASE",
                                                         "https://api.deepseek.com")
    model = cfg.get("llm_model") or os.environ.get("LLM_MODEL", "deepseek-chat")
    # 领域优先取任务自身选择的（job.domain），否则用全局设置
    domain = domain or cfg.get("domain", "religion")
    if not api_key:
        raise RuntimeError("未配置 LLM API key（设置页可填，或设 LLM_API_KEY 环境变量）")

    if hotwords is None:
        # 读该领域的热词表（正确词列表，供 LLM 参考）
        from pipeline import _load_hotword_map
        pairs = _load_hotword_map(domain)
        hotwords = [right for _, right in pairs]

    sentences = result["sentences"]
    # 按时间窗口切上下文段
    groups, cur, w0 = [], [], None
    for s in sentences:
        if w0 is None or s["start"] - w0 >= WINDOW_SEC:
            if cur:
                groups.append(cur)
            cur, w0 = [s], s["start"]
        else:
            cur.append(s)
    if cur:
        groups.append(cur)

    all_out = []
    for i, grp in enumerate(groups):
        if cb:
            cb(95 + 5 * i / len(groups), f"LLM 纠错 {i+1}/{len(groups)} 段...")
        seg_text = "\n".join(
            f"讲话人{s['spk_display']}  {fmt_mmss(s['start'])}  {s['text']}" for s in grp)
        msgs = build_messages(seg_text, domain, hotwords)
        try:
            resp = call_llm(api_base, api_key, model, msgs)
            fixed = json.loads(resp).get("corrected_text", "")
            lines = [l.strip() for l in fixed.splitlines() if l.strip()]
        except Exception:
            lines = seg_text.splitlines()

        # 按行回填，行数不一致则保留原文
        import re
        if len(lines) == len(grp):
            pat = re.compile(r"^讲话人\d+\s+\d[\d:]*\s+(.*)$")
            for s, line in zip(grp, lines):
                m = pat.match(line)
                s["text"] = (m.group(1).strip() if m else line) or s["text"]
        all_out.extend(grp)

    # 写 polished 文件
    result_polished = dict(result)
    result_polished["sentences"] = all_out
    result_polished["polished"] = True
    (Path(job_dir) / "result_polished.json").write_text(
        json.dumps(result_polished, ensure_ascii=False, indent=2))

    with open(Path(job_dir) / "result_polished.txt", "w", encoding="utf-8") as f:
        for s in all_out:
            f.write(f"讲话人{s['spk_display']}  {fmt_mmss(s['start'])}  {s['text']}\n")

    # polished docx：复用 pipeline 的 docx 生成（简单重写一份）
    _write_docx(job_dir, result_polished)
    return result_polished


def _write_docx(job_dir, result):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    h = doc.add_heading(result["title"] + "（已校对）", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cur_spk, buf, t_start = None, [], None
    def flush():
        nonlocal cur_spk, buf, t_start
        if cur_spk is not None and buf:
            p = doc.add_paragraph()
            lab = p.add_run(f"讲话人{cur_spk}  [{fmt_mmss(t_start)}]  ")
            lab.bold = True
            lab.font.color.rgb = RGBColor(0x1a, 0x5c, 0xb0)
            p.add_run("".join(buf))
        cur_spk, buf, t_start = None, [], None

    for s in result["sentences"]:
        if s["spk_display"] != cur_spk:
            flush()
            cur_spk, t_start = s["spk_display"], s["start"]
        buf.append(s["text"])
    flush()
    doc.save(str(Path(job_dir) / "result_polished.docx"))
